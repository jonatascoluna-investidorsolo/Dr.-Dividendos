import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def import_sheet(path: Path):
    wb = load_workbook(path, data_only=False, read_only=True)
    ws = wb["Dados"]
    headers = [ws.cell(7, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(8, ws.max_row + 1):
        ticker = ws.cell(r, 2).value
        company = ws.cell(r, 1).value
        if not ticker and not company:
            continue
        item = {}
        for c, header in enumerate(headers, start=1):
            if header:
                item[str(header)] = clean(ws.cell(r, c).value)
        rows.append(item)
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("excel")
    parser.add_argument("--output", default="data/reference/dados.json")
    args = parser.parse_args()
    rows = import_sheet(Path(args.excel))
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Importadas {len(rows)} linhas de Dados -> {output}")


if __name__ == "__main__":
    main()
