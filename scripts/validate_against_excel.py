from pathlib import Path
import json
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from dd_engine.models import CompanyInput
from dd_engine.service import calculate_company
from dd_engine.settings import MethodologyParameters

EXCEL = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Planilha_Doutor_dos_Dividendos_v2.xlsx")
wb = load_workbook(EXCEL, data_only=True, read_only=False)
ws = wb["Dados"]
panel = wb["Painel"]
params = MethodologyParameters(
    bazin_min_yield=panel["C15"].value,
    gordon_required_return=panel["C16"].value,
    max_growth=panel["C17"].value,
    max_net_debt_ebitda=panel["C18"].value,
    projection_years=int(panel["C19"].value),
    projective_required_yield=panel["C20"].value,
    max_healthy_payout=panel["C21"].value,
)
outs = [
    ("O", "lpa"), ("P", "dpa"), ("Q", "dividend_yield"), ("R", "pe_current"),
    ("S", "pb_current"), ("T", "graham_value"), ("U", "graham_upside"),
    ("V", "bazin_ceiling"), ("W", "bazin_margin"), ("X", "applied_growth"),
    ("Y", "gordon_value"), ("Z", "gordon_upside"), ("AA", "projected_lpa"),
    ("AB", "projected_dpa"), ("AC", "projective_ceiling"), ("AD", "projective_upside"),
    ("AE", "dy_score"), ("AF", "margin_score"), ("AG", "debt_score"),
    ("AH", "growth_score"), ("AI", "payout_score"), ("AJ", "quality_score"),
    ("AK", "seal"),
]

def close(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if a is None or b is None:
        return False
    if isinstance(a, str) or isinstance(b, str):
        return str(a) == str(b)
    return abs(float(a) - float(b)) <= 1e-6 * max(1, abs(float(a)), abs(float(b)))

companies = []
for r in range(8, 108):
    vals = [ws.cell(r, c).value for c in range(1, 12)]
    if not vals[1]:
        continue
    inp = CompanyInput(
        ticker=str(vals[1]), company_name=str(vals[0] or ""), sector=vals[2],
        shares=vals[3], projected_net_income=vals[4], book_value_per_share=vals[5],
        net_debt_ebitda=vals[6], earnings_cagr_5y=vals[7], payout_expected=vals[9],
        current_price=vals[10],
    )
    result = calculate_company(inp, params)
    diffs = []
    for col, attr in outs:
        excel_value = ws[f"{col}{r}"].value
        engine_value = getattr(result, attr)
        if not close(excel_value, engine_value):
            diffs.append({"cell": f"{col}{r}", "metric": attr, "excel": excel_value, "engine": engine_value})
    companies.append({"row": r, "ticker": vals[1], "company": vals[0], "differences": diffs})

result = {
    "excel": str(EXCEL),
    "companies": len(companies),
    "companies_with_differences": sum(bool(x["differences"]) for x in companies),
    "differences": [x for x in companies if x["differences"]],
}
out = ROOT / "validation_report.json"
out.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str))
print(json.dumps({k: result[k] for k in ("companies", "companies_with_differences")}, ensure_ascii=False))
if result["companies_with_differences"]:
    for item in result["differences"][:10]:
        print(item)
    raise SystemExit(1)
print("VALIDATION_OK: motor reproduces cached Excel outputs for all companies and metrics tested.")
