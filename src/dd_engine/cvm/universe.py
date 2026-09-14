from __future__ import annotations
from dataclasses import dataclass, asdict
from pathlib import Path
import csv
from openpyxl import load_workbook

@dataclass(frozen=True)
class UniverseItem:
    excel_row: int
    company_name: str
    ticker: str
    likely_bdr: bool
    sector: str | None

def load_excel_universe(path: str | Path) -> list[UniverseItem]:
    wb=load_workbook(path, data_only=False, read_only=True)
    ws=wb['Dados']
    out=[]
    for r in range(8, ws.max_row+1):
        ticker=ws.cell(r,2).value
        if not ticker: continue
        ticker=str(ticker).strip().upper()
        out.append(UniverseItem(
            excel_row=r,
            company_name=str(ws.cell(r,1).value or '').strip(),
            ticker=ticker,
            likely_bdr=ticker.endswith(('34','35','39')),
            sector=str(ws.cell(r,3).value or '').strip() or None,
        ))
    return out

def write_manifest(items: list[UniverseItem], path: str | Path) -> None:
    p=Path(path); p.parent.mkdir(parents=True, exist_ok=True)
    with p.open('w', newline='', encoding='utf-8') as f:
        w=csv.DictWriter(f, fieldnames=list(asdict(items[0]).keys()))
        w.writeheader()
        for item in items: w.writerow(asdict(item))
